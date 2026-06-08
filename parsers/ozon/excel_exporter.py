# core/excel_exporter.py by the88788.com
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core.utils import stamp_for_filename, sanitize_rows_for_excel


def export_snapshot(
    base_dir: Path,
    data_rows: list[dict],
    relation_rows: list[dict] | None = None,
    error_rows: list[dict] | None = None,
    fixed_path: Path | None = None,
) -> str:
    output_dir = base_dir / 'output'
    output_dir.mkdir(parents=True, exist_ok=True)
    output_file = fixed_path if fixed_path is not None else output_dir / f'ozon_parse_{stamp_for_filename()}.xlsx'
    safe_data_rows = sanitize_rows_for_excel(data_rows)
    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        pd.DataFrame(safe_data_rows).to_excel(writer, index=False, sheet_name='Данные')
        if error_rows:
            pd.DataFrame(sanitize_rows_for_excel(error_rows)).to_excel(
                writer, index=False, sheet_name='Ошибки'
            )

    _style_workbook(output_file)
    return str(output_file)


def _style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill(fill_type='solid', fgColor='D9EAF7')
    for ws in wb.worksheets:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for column_cells in ws.columns:
            values = [str(c.value) if c.value is not None else '' for c in column_cells]
            width = min(max(len(v) for v in values) + 2, 45)
            ws.column_dimensions[column_cells[0].column_letter].width = width
    wb.save(path)