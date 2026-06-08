# core/excel_exporter.py by the88788.com
from __future__ import annotations

from pathlib import Path

import pandas as pd
from openpyxl import load_workbook
from openpyxl.styles import Alignment, Font, PatternFill

from core.utils import stamp_for_filename, sanitize_rows_for_excel


def _compute_spp_price(price_str: str, spp_percent: int) -> str:
    """Вычисляет цену с СПП: price_current * (1 - spp_percent/100), округляет до целых."""
    try:
        clean = ''.join(c for c in str(price_str) if c.isdigit())
        if not clean:
            return ''
        val = float(clean)
        if val <= 0:
            return ''
        return str(round(val * (1 - spp_percent / 100)))
    except (ValueError, TypeError):
        return ''


def export_snapshot(
    base_dir: Path,
    data_rows: list[dict],
    relation_rows: list[dict] | None = None,
    error_rows: list[dict] | None = None,
    fixed_path: Path | None = None,
    spp_percent: int = 25,
) -> str:
    output_dir = base_dir / 'output'
    output_dir.mkdir(parents=True, exist_ok=True)
    output_file = fixed_path if fixed_path is not None else output_dir / f'wb_parse_{stamp_for_filename()}.xlsx'
    safe_data_rows = sanitize_rows_for_excel(data_rows)

    # Вставляем price_spp сразу после price_wallet
    for row in safe_data_rows:
        row['price_spp'] = _compute_spp_price(row.get('price_current', ''), spp_percent)

    # Упорядочиваем колонки: price_spp после price_wallet
    if safe_data_rows:
        cols = list(safe_data_rows[0].keys())
        if 'price_spp' in cols and 'price_wallet' in cols:
            cols.remove('price_spp')
            wallet_idx = cols.index('price_wallet')
            cols.insert(wallet_idx + 1, 'price_spp')
            safe_data_rows = [{c: row.get(c, '') for c in cols} for row in safe_data_rows]

    with pd.ExcelWriter(output_file, engine='openpyxl') as writer:
        pd.DataFrame(safe_data_rows).to_excel(writer, index=False, sheet_name='Данные')
        if error_rows:
            pd.DataFrame(sanitize_rows_for_excel(error_rows)).to_excel(
                writer, index=False, sheet_name='Ошибки'
            )

    _style_workbook(output_file)
    return str(output_file)


def _style_workbook(path: Path) -> None:
    wb = load_workbook(path)
    header_fill = PatternFill(fill_type='solid', fgColor='D9EAF7')
    for ws in wb.worksheets:
        ws.freeze_panes = 'A2'
        ws.auto_filter.ref = ws.dimensions
        for cell in ws[1]:
            cell.font = Font(bold=True)
            cell.fill = header_fill
            cell.alignment = Alignment(horizontal='center', vertical='center', wrap_text=True)
        for column_cells in ws.columns:
            values = [str(c.value) if c.value is not None else '' for c in column_cells]
            width = min(max(len(v) for v in values) + 2, 45)
            ws.column_dimensions[column_cells[0].column_letter].width = width
    wb.save(path)
